"""
Document generation tool — real PDF and PPTX creation.

Uses reportlab for PDF and python-pptx for PPTX. No LLM calls happen here;
this tool takes already-decided content (title + list of sections/slides)
and lays it out. Content generation is the caller's job.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def generate_pdf(output_path: str | Path, title: str, sections: list[tuple[str, str]]) -> Path:
    """sections: list of (heading, body_text) pairs."""
    output_path = Path(output_path)
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(output_path), pagesize=letter)
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for heading, body in sections:
        story.append(Paragraph(heading, styles["Heading2"]))
        story.append(Paragraph(body, styles["BodyText"]))
        story.append(Spacer(1, 10))
    doc.build(story)
    return output_path


def generate_pptx(output_path: str | Path, title: str, slides: list[tuple[str, str]]) -> Path:
    """slides: list of (title, bullet_text) pairs, one per slide."""
    output_path = Path(output_path)
    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    first = prs.slides.add_slide(title_layout)
    first.shapes.title.text = title

    for slide_title, body in slides:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = slide_title
        body_placeholder = slide.placeholders[1]
        body_placeholder.text = body

    prs.save(str(output_path))
    return output_path
